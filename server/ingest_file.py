#!/usr/bin/env python3
"""
ingest_file.py — turn a file into memories (memanto's `upload`, the local way).

It extracts text, splits it into deterministic chunks, and stores each chunk as a
memory tagged with the filename and marked provenance origin=imported,
source=file:<name>#chunk<i>. There is NO LLM and no summarization: chunking is
pure rule-based (paragraph boundaries with a size target + slight overlap), and
only the local embedder runs (to store the chunks).

Text formats (.txt/.md/.csv/.tsv/.json/.log) need only the stdlib. PDF/DOCX/XLSX
need optional parsers, isolated in requirements-ingest.txt so they are installed
only by people who ingest those formats:
    .venv/bin/pip install -r requirements-ingest.txt

Usage -- writing requires exclusive store access, so STOP THE BACKEND first
(same rule as the migration scripts):
    launchctl kill TERM gui/$(id -u)/com.mem0mcp.server   # or close all clients
    .venv/bin/python server/ingest_file.py notes.md
    .venv/bin/python server/ingest_file.py report.pdf --target-chars 1000 --overlap 120
    .venv/bin/python server/ingest_file.py notes.md --dry-run   # preview chunks, write nothing
"""
import os
import re
import sys
import json
import argparse

from mem0_store import is_backend_up

_TEXT_EXTS = {".txt", ".text", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv"}


def _extract_pdf(path: str) -> str:
    try:
        import pypdf
    except ImportError as e:
        raise ImportError("PDF ingest needs pypdf — install the optional parsers:\n"
                          "  .venv/bin/pip install -r requirements-ingest.txt") from e
    reader = pypdf.PdfReader(path)
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(path: str) -> str:
    try:
        import docx
    except ImportError as e:
        raise ImportError("DOCX ingest needs python-docx — install the optional parsers:\n"
                          "  .venv/bin/pip install -r requirements-ingest.txt") from e
    return "\n\n".join(p.text for p in docx.Document(path).paragraphs if p.text)


def _extract_xlsx(path: str) -> str:
    try:
        import openpyxl
    except ImportError as e:
        raise ImportError("XLSX ingest needs openpyxl — install the optional parsers:\n"
                          "  .venv/bin/pip install -r requirements-ingest.txt") from e
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"# {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                out.append("\t".join(cells))
    return "\n".join(out)


def extract_text(path: str) -> str:
    """Extract plain text from a file. Stdlib for text/JSON/CSV; optional parsers
    (pypdf/python-docx/openpyxl) for .pdf/.docx/.xlsx with a clear install hint if
    missing. Deterministic; no network, no LLM."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".json":
        with open(path, encoding="utf-8") as f:
            raw = f.read()
        try:
            return json.dumps(json.loads(raw), ensure_ascii=False, indent=2)
        except ValueError:
            return raw
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        return _extract_docx(path)
    if ext in (".xlsx", ".xlsm"):
        return _extract_xlsx(path)
    if ext in _TEXT_EXTS or ext == "":
        with open(path, encoding="utf-8") as f:
            return f.read()
    try:
        with open(path, encoding="utf-8") as f:
            return f.read()
    except (UnicodeDecodeError, OSError) as e:
        raise ValueError(f"unsupported or unreadable file '{path}': {e}") from e


def _split_long(para: str, target: int, overlap: int) -> list:
    """Hard-split a single over-long paragraph into <=target windows that overlap
    by `overlap` chars (so a fact split across the boundary is still recoverable)."""
    if len(para) <= target:
        return [para]
    step = max(1, target - overlap)
    return [para[i:i + target] for i in range(0, len(para), step)]


def chunk_text(text: str, target_chars: int = 800, overlap: int = 100) -> list:
    """Split text into deterministic chunks. Paragraphs (blank-line separated) are
    greedily packed up to ~target_chars; a paragraph longer than the target is
    hard-split with `overlap` chars of carry-over. Pure + deterministic: same input
    always yields the same chunks. Returns [] for empty input."""
    text = (text or "").strip()
    if not text:
        return []
    target_chars = max(1, target_chars)
    overlap = max(0, min(overlap, target_chars - 1))
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    units = []
    for p in paras:
        units.extend(_split_long(p, target_chars, overlap))
    chunks, cur = [], ""
    for u in units:
        if cur and len(cur) + 2 + len(u) > target_chars:
            chunks.append(cur)
            cur = u
        else:
            cur = (cur + "\n\n" + u) if cur else u
    if cur:
        chunks.append(cur)
    return chunks


def ingest(path: str, user: str = "", target_chars: int = 800, overlap: int = 100) -> dict:
    """Extract -> chunk -> store every chunk as a memory (origin=imported,
    source=file:<name>#chunk<i>, tag=<filename stem>). Stores in ONE locked pass via
    the server's batch path. Returns {path, chunks, ids, results}. The CALLER must
    ensure exclusive store access (the CLI refuses while the backend is up)."""
    chunks = chunk_text(extract_text(path), target_chars, overlap)
    base = os.path.basename(path)
    if not chunks:
        return {"path": path, "chunks": 0, "ids": [], "results": []}
    tag = os.path.splitext(base)[0].replace(" ", "-")
    items = [{"text": c, "origin": "imported", "source": f"file:{base}#chunk{i + 1}",
              "tags": tag} for i, c in enumerate(chunks)]
    # Deferred: importing the server constructs the embedder + opens Chroma. Done
    # only when actually writing, AFTER the backend-up guard in main().
    import mem0_mcp_server as srv
    results = srv._add_many(items, user or srv.DEFAULT_USER)
    return {"path": path, "chunks": len(chunks),
            "ids": [r["id"] for r in results if r.get("id")], "results": results}


def main():
    ap = argparse.ArgumentParser(description="Ingest a file into mem0 memories (chunked).")
    ap.add_argument("path", help="file to ingest (.txt/.md/.csv/.json/.log, or .pdf/.docx/.xlsx)")
    ap.add_argument("--user", default="", help="user_id (default: MEM0_DEFAULT_USER)")
    ap.add_argument("--target-chars", type=int, default=800, help="approx chunk size (default 800)")
    ap.add_argument("--overlap", type=int, default=100, help="chunk overlap chars (default 100)")
    ap.add_argument("--dry-run", action="store_true", help="preview chunks; write nothing")
    args = ap.parse_args()

    if not os.path.isfile(args.path):
        sys.exit(f"no such file: {args.path}")

    if args.dry_run:
        chunks = chunk_text(extract_text(args.path), args.target_chars, args.overlap)
        print(f"[dry-run] {args.path}: {len(chunks)} chunk(s) "
              f"(~{args.target_chars} chars, overlap {args.overlap}); nothing written.")
        for i, c in enumerate(chunks, 1):
            preview = c[:200].replace("\n", " ")
            print(f"  chunk {i} ({len(c)} chars): {preview}{'…' if len(c) > 200 else ''}")
        return

    # Writing requires exclusive store access -- refuse while the backend is up
    # (concurrent Chroma writers risk corruption; same rule as the migrations).
    host = os.environ.get("MEM0_MCP_HOST", "127.0.0.1")
    port = int(os.environ.get("MEM0_MCP_PORT", "8765"))
    if is_backend_up(host, port):
        sys.exit(f"Backend is running on {host}:{port}. Stop it first (or use --dry-run):\n"
                 f"  launchctl kill TERM gui/$(id -u)/com.mem0mcp.server")

    res = ingest(args.path, user=args.user, target_chars=args.target_chars, overlap=args.overlap)
    if not res["chunks"]:
        print(f"no text extracted from {args.path}; nothing ingested.")
        return
    print(f"✅ Ingested {res['chunks']} chunk(s) from {args.path} -> {len(res['ids'])} memories "
          f"(origin=imported, tag=#{os.path.splitext(os.path.basename(args.path))[0].replace(' ', '-')}).")


if __name__ == "__main__":
    main()
